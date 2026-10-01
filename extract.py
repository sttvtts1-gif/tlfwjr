"""RPS프로젝트 엑셀(연도별) -> data.js 생성. 사용: python3 extract.py 2025.xlsx 2026.xlsx > ../data.js"""
import openpyxl,json,datetime,re,sys
types={'발전사업(A)':'발전사업','자가소비(B)':'자가소비','PPA 사업(C)':'PPA','리스사업(D)':'리스','임대사업(E)':'임대','환경부사업(F)':'환경부','프로젝트(G)':'프로젝트','협력사 EPC(H)':'협력사EPC','협력사 하도(I)':'협력사하도','기타(Z)':'기타(Z)'}
branch={'안성지사':['경기','강원','충남','충북','서울','인천','대전','세종'],'부산지사':['경남','경북','전남','전북','부산','대구','울산','제주','광주']}
b_of={r:b for b,rs in branch.items() for r in rs}
def s(v): return (str(v).strip() if v is not None else '')
def dt(v):
    if isinstance(v,datetime.datetime): return v.strftime('%Y-%m-%d')
    m=re.match(r'^(\d{2,4})[.\-/](\d{1,2})[.\-/](\d{1,2})',s(v))
    if m: y,mo,d=m.groups(); y=int(y); y=y+2000 if y<100 else y; return f'{y:04d}-{int(mo):02d}-{int(d):02d}'
    return s(v)
rows=[];targets={}
for f in sys.argv[1:]:
    yr=int(re.search(r'(20\d\d)',f).group(1))
    wb=openpyxl.load_workbook(f,read_only=True,data_only=True)
    t={}
    for r in wb['정리base'].iter_rows(min_row=3,max_row=10,values_only=True):
        if r[3] in ('직영','외주','입찰') and r[3] not in t: t[r[3]]=[float(x or 0) for x in r[4:16]]
    targets[yr]=t
    for sh,tn in types.items():
        if sh not in wb.sheetnames: continue
        ws=wb[sh]; hdr=[s(c).replace('\n','') for c in next(ws.iter_rows(min_row=2,max_row=2,values_only=True))]
        ix={}
        for i,h in enumerate(hdr):
            if h and h not in ix: ix[h]=i
        for r in ws.iter_rows(min_row=3,values_only=True):
            no=r[ix['NO.']]; cap=r[ix['계약용량(kWp)']]; mon=r[ix['월']]
            if not isinstance(no,(int,float)) or not isinstance(cap,(int,float)) or not isinstance(mon,(int,float)): continue
            reg=s(r[ix['지역']]); amt=r[ix['공사금액(부가세제외)']]
            rows.append(dict(y=yr,m=int(mon),t=tn,ch=re.sub(r'\(.\)','',s(r[ix['유입유형']])),d=dt(r[ix['계약일']]),cap=round(float(cap),3),
              own=s(r[ix['담당자']]),cons=s(r[ix['컨설팅']]),name=s(r[ix['발전소명']]),reg=reg,br=b_of.get(reg,'기타'),
              inflow='소개' if s(r[ix['유입']])=='소개' else '직접',loc=s(r[ix['위치']]),
              inst=s(r[ix['설치유형']]).replace('기존시설물(A)','기존시설물').replace('일반부지(B)','일반부지'),
              amt=float(amt) if isinstance(amt,(int,float)) else 0))
print('window.GS_DATA='+json.dumps({'rows':rows,'targets':targets,'branch':branch},ensure_ascii=False,separators=(',',':'))+';')
